#!/usr/bin/env python3
"""Parse PHIẾU CHỈ ĐỊNH (PK gửi) → 2 danh mục chuẩn hoá: dịch vụ/CLS + thuốc.

Nguồn: scripts/catalog_src/PHIEU_CHI_DINH_update.docx (python-docx).
  - Table 0 + Table 1 = dịch vụ/CLS theo nhóm-cột (Tầng / Thai / Nội tiết-phụ
    khoa / Thủ thuật / Chụp phim).
  - Table 2 = thuốc, mỗi dòng (row) = 1 list ngăn cách bởi ',' và ';',
    biến thể nằm trong ngoặc.

Xuất scripts/catalog_out/{services.csv, drugs.csv} + in thống kê đối chiếu.

BOUNDARY: giữ name_raw VERBATIM; tách variant best-effort; dòng không chắc →
needs_review=TRUE, KHÔNG bỏ sót. KHÔNG bịa giá (giá nằm ở DB, NULL).
"""
from __future__ import annotations

import csv
import os
import re

import docx

SRC = "scripts/catalog_src/PHIEU_CHI_DINH_update.docx"
OUT_DIR = "scripts/catalog_out"


def clean(s: str) -> str:
    """nbsp→space, gộp khoảng trắng, trim. KHÔNG đổi định danh."""
    return re.sub(r"\s+", " ", s.replace("\xa0", " ")).strip()


def is_floor(s: str) -> bool:
    return s.startswith("Tầng")


# --------------------------------------------------------------------------- #
# DỊCH VỤ / CLS (Table 0, Table 1)
# --------------------------------------------------------------------------- #
def parse_services(doc: "docx.document.Document") -> tuple[list[dict], int]:
    services: list[dict] = []
    for t in doc.tables:
        if len(t.columns) < 2:  # bảng thuốc 1 cột → bỏ
            continue
        header = [clean(c.text) for c in t.rows[0].cells]
        # gom các cột liền kề cùng header thành 1 "nhóm"; cột đầu nhóm = cột tên
        groups: list[tuple[int, str]] = []
        c, n = 0, len(header)
        while c < n:
            label, start = header[c], c
            while c + 1 < n and header[c + 1] == label:
                c += 1
            groups.append((start, label))
            c += 1
        for start, label in groups:
            cur_tang = label if is_floor(label) else ""
            for ri in range(1, len(t.rows)):
                cell = clean(t.rows[ri].cells[start].text)
                if not cell or cell == label:
                    continue
                if is_floor(cell):  # ô "Tầng 2/ Tầng 4" = ngữ cảnh tầng, không phải dịch vụ
                    cur_tang = cell
                    continue
                services.append(
                    {"name": cell, "group": label, "tang": cur_tang, "name_raw": cell}
                )
    # dedupe (group, name) cho khớp unique index service_price(group, service_code)
    seen: set[tuple[str, str]] = set()
    uniq: list[dict] = []
    dups = 0
    for s in services:
        k = (s["group"], s["name"])
        if k in seen:
            dups += 1
            continue
        seen.add(k)
        uniq.append(s)
    return uniq, dups


# --------------------------------------------------------------------------- #
# THUỐC (Table 2)
# --------------------------------------------------------------------------- #
def split_drug_line(line: str) -> list[str]:
    """Tách ',' và ';' nhưng KHÔNG tách khi đang trong ngoặc (giữ 'Letrozole (10v, 15v)')."""
    toks, buf, depth = [], "", 0
    for ch in line:
        if ch == "(":
            depth += 1
            buf += ch
        elif ch == ")":
            depth = max(0, depth - 1)
            buf += ch
        elif ch in ",;" and depth == 0:
            if buf.strip():
                toks.append(buf)
            buf = ""
        else:
            buf += ch
    if buf.strip():
        toks.append(buf)
    return toks


def parse_drug_token(tok: str, line_no: int) -> dict:
    raw = clean(tok)
    parens = re.findall(r"\(([^()]*)\)", raw)
    base = clean(re.sub(r"\([^()]*\)", "", raw)).strip(" :,-")
    variant = "; ".join(p.strip() for p in parens if p.strip())
    needs = False
    if ":" in raw:  # cú pháp lạ (vd "Utrogestan (Đ) (1v/2v): (U)")
        needs = True
    if len(parens) > 1:  # nhiều cụm ngoặc → biến thể không chắc
        needs = True
    if "/" in base:  # slash NGOÀI ngoặc → synonyms / liều cần dược xác nhận
        needs = True
    if not base:
        base = raw
        needs = True
    return {
        "name_base": base,
        "name_raw": raw,
        "variant": variant,
        "group_line": f"L{line_no}",
        "needs_review": needs,
    }


def parse_drugs(doc: "docx.document.Document") -> list[dict]:
    drugs: list[dict] = []
    drug_table = next((t for t in doc.tables if len(t.columns) == 1), None)
    if drug_table is None:
        return drugs
    for ri, row in enumerate(drug_table.rows, start=1):
        line = clean(row.cells[0].text)
        if not line:
            continue
        for tok in split_drug_line(line):
            drugs.append(parse_drug_token(tok, ri))
    return drugs


def main() -> None:
    os.makedirs(OUT_DIR, exist_ok=True)
    doc = docx.Document(SRC)

    services, svc_dups = parse_services(doc)
    drugs = parse_drugs(doc)

    with open(os.path.join(OUT_DIR, "services.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["name", "group", "tang", "name_raw"])
        w.writeheader()
        w.writerows(services)

    with open(os.path.join(OUT_DIR, "drugs.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(
            f, fieldnames=["name_base", "name_raw", "variant", "group_line", "needs_review"]
        )
        w.writeheader()
        for d in drugs:
            row = dict(d)
            row["needs_review"] = "TRUE" if d["needs_review"] else "FALSE"
            w.writerow(row)

    nr = sum(1 for d in drugs if d["needs_review"])
    print("=== SERVICES ===")
    by_group: dict[str, int] = {}
    for s in services:
        by_group[s["group"]] = by_group.get(s["group"], 0) + 1
    for g, c in by_group.items():
        print(f"  {g}: {c}")
    print(f"  TỔNG dịch vụ (sau dedupe): {len(services)}  (đã gộp {svc_dups} dòng trùng)")
    print("\n=== DRUGS ===")
    by_line: dict[str, int] = {}
    for d in drugs:
        by_line[d["group_line"]] = by_line.get(d["group_line"], 0) + 1
    for ln, c in sorted(by_line.items(), key=lambda kv: int(kv[0][1:])):
        print(f"  {ln}: {c}")
    print(f"  TỔNG thuốc: {len(drugs)}  |  needs_review=TRUE: {nr}")


if __name__ == "__main__":
    main()
